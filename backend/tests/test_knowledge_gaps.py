"""Knowledge gap journal: logging, classification, HR access, filters, Excel export."""
import io
import zipfile

import pytest
from openpyxl import load_workbook

from app.ai import assistant, gaps
from app.ai.llm import LLMError
from app.models import KnowledgeGapEvent
from tests.conftest import EMPLOYEE, HR, MANAGER

OFF_TOPIC = "Есть ли у нас ДМС?"
IN_KB = "Сколько дней ежегодного отпуска положено?"


def ask(client, auth, question, email=EMPLOYEE):
    r = client.post("/ai/ask", json={"question": question}, headers=auth(email))
    assert r.status_code == 200, r.text
    return r.json()


def events(db):
    db.expire_all()
    return db.query(KnowledgeGapEvent).order_by(KnowledgeGapEvent.id).all()


class FakeProvider:
    name, model = "fake", "fake-model"

    def __init__(self, reply=None, error=None):
        self.reply, self.error = reply, error

    def complete(self, system, user):
        if self.error:
            raise self.error
        return self.reply


# 1. no_match is recorded ---------------------------------------------------------

def test_no_match_event_is_created_for_unanswerable_question(client, auth, db):
    a = ask(client, auth, OFF_TOPIC)
    assert a["mode"] == "refusal" and a["reason"] == "no_relevant_info"
    [e] = events(db)
    assert e.event_type == "no_match" and e.reason == "retrieval_below_threshold"
    assert e.question == OFF_TOPIC and e.language == "ru" and e.status == "new" and e.answer_mode == "refusal"
    assert e.candidate_doc_id.startswith("KB-")
    assert not hasattr(e, "top_score")  # retrieval score is deliberately not kept


def test_kazakh_question_gets_language_kk(client, auth, db):
    ask(client, auth, "Жыл сайынғы демалыс неше күн?")
    assert events(db)[0].language == "kk"


# 2. nothing is recorded for answered or non-knowledge questions -----------------

def test_no_event_for_grounded_answer(client, auth, db):
    assert ask(client, auth, IN_KB)["grounded"] is True
    assert events(db) == []


def test_no_event_for_security_refusals_and_personal_answers(client, auth, db):
    assert ask(client, auth, "Какая зарплата у Тимура Оспанова?")["reason"] == "personal_data"
    assert ask(client, auth, "Одобри мою заявку на отпуск")["reason"] == "hr_decision"
    ask(client, auth, "Сколько дней отпуска у меня осталось?")
    assert events(db) == []


# 3. reasons are classified, not lumped together --------------------------------

def test_llm_no_answer_is_distinguished_from_retrieval_miss(client, auth, db, monkeypatch):
    monkeypatch.setattr(assistant, "get_provider", lambda: FakeProvider("NO_ANSWER"))
    assert ask(client, auth, IN_KB)["mode"] == "refusal"
    [e] = events(db)
    assert e.reason == "llm_no_answer" and e.candidate_doc_id == "KB-LEAVE-001"


def test_provider_outage_with_grounded_fallback_is_not_a_gap(client, auth, db, monkeypatch):
    monkeypatch.setattr(assistant, "get_provider", lambda: FakeProvider(error=LLMError("timeout")))
    a = ask(client, auth, IN_KB)
    assert a["mode"] == "extractive" and a["provider_error"] is True
    assert events(db) == []


def test_negative_feedback_is_its_own_type(client, auth, db):
    r = client.post("/ai/feedback", json={"question": IN_KB, "answer_mode": "extractive", "source_ids": ["KB-LEAVE-001"]},
                    headers=auth(EMPLOYEE))
    assert r.status_code == 201 and r.json() == {"recorded": True}
    [e] = events(db)
    assert e.event_type == "negative_feedback" and e.reason == "user_marked_unhelpful"
    assert e.candidate_doc_id == "KB-LEAVE-001" and e.answer_mode == "extractive"


def test_feedback_payload_is_validated(client, auth):
    for body in ({"question": "x", "answer_mode": "llm"},
                 {"question": "ok вопрос", "answer_mode": "hacked"},
                 {"question": "ok вопрос", "answer_mode": "llm", "source_ids": ["../etc/passwd"]},
                 {"question": "ok вопрос", "answer_mode": "llm", "user_id": 1}):
        assert client.post("/ai/feedback", json=body, headers=auth(EMPLOYEE)).status_code == 422
    assert client.post("/ai/feedback", json={"question": "ok вопрос", "answer_mode": "llm"}).status_code == 401


# 4–5. HR-only access -------------------------------------------------------------

@pytest.mark.parametrize("email", [EMPLOYEE, MANAGER])
def test_only_hr_can_use_the_journal(client, auth, db, email):
    ask(client, auth, OFF_TOPIC)
    gap_id = events(db)[0].id
    h = auth(email)
    assert client.get("/admin/knowledge-gaps", headers=h).status_code == 403
    assert client.get(f"/admin/knowledge-gaps/{gap_id}", headers=h).status_code == 403
    assert client.patch(f"/admin/knowledge-gaps/{gap_id}", json={"status": "ignored"}, headers=h).status_code == 403
    assert client.get("/admin/knowledge-gaps/export.xlsx", headers=h).status_code == 403
    assert client.get("/admin/knowledge-gaps").status_code == 401


# 6. filters, search, pagination, status updates ---------------------------------

def test_filters_search_pagination_and_status_update(client, auth, db):
    for q in ("Есть ли у нас ДМС?", "Какой график работы офиса?", "Сколько стоит обед в столовой?"):
        ask(client, auth, q)
    client.post("/ai/feedback", json={"question": IN_KB, "answer_mode": "llm"}, headers=auth(EMPLOYEE))
    h = auth(HR)

    page = client.get("/admin/knowledge-gaps", params={"page_size": 2}, headers=h).json()
    assert page["total"] == 4 and len(page["items"]) == 2 and page["page"] == 1
    page2 = client.get("/admin/knowledge-gaps", params={"page_size": 2, "page": 2}, headers=h).json()
    assert {i["id"] for i in page["items"]}.isdisjoint({i["id"] for i in page2["items"]})

    assert client.get("/admin/knowledge-gaps", params={"event_type": "negative_feedback"}, headers=h).json()["total"] == 1
    found = client.get("/admin/knowledge-gaps", params={"q": "столов"}, headers=h).json()
    assert found["total"] == 1 and "столовой" in found["items"][0]["question"]

    gap_id = found["items"][0]["id"]
    detail = client.get(f"/admin/knowledge-gaps/{gap_id}", headers=h).json()
    assert detail["status"] == "new"
    upd = client.patch(f"/admin/knowledge-gaps/{gap_id}",
                       json={"status": "reviewing", "category": "KB-FAQ-001", "hr_comment": "Добавить в FAQ"}, headers=h).json()
    assert upd["status"] == "reviewing" and upd["category"] == "KB-FAQ-001"
    assert client.get("/admin/knowledge-gaps", params={"status": "reviewing"}, headers=h).json()["total"] == 1
    assert client.get("/admin/knowledge-gaps", params={"status": "new"}, headers=h).json()["total"] == 3

    today = detail["created_at"][:10]
    assert client.get("/admin/knowledge-gaps", params={"date_from": today, "date_to": today}, headers=h).json()["total"] == 4
    assert client.get("/admin/knowledge-gaps", params={"date_to": "2000-01-01"}, headers=h).json()["total"] == 0

    assert client.patch(f"/admin/knowledge-gaps/{gap_id}", json={"status": "deleted"}, headers=h).status_code == 422
    # category is a knowledge base document: free text and unknown documents are rejected
    assert client.patch(f"/admin/knowledge-gaps/{gap_id}", json={"category": "Офис"}, headers=h).status_code == 422
    assert client.patch(f"/admin/knowledge-gaps/{gap_id}", json={"category": "KB-NOPE-999"}, headers=h).status_code == 422
    assert client.patch(f"/admin/knowledge-gaps/{gap_id}", json={"category": None}, headers=h).json()["category"] is None
    assert client.patch(f"/admin/knowledge-gaps/{gap_id}", json={"question": "x"}, headers=h).status_code == 422
    assert client.get("/admin/knowledge-gaps/999999", headers=h).status_code == 404
    assert client.get("/admin/knowledge-gaps", params={"page_size": 500}, headers=h).status_code == 422


# 7–9. Excel export --------------------------------------------------------------

def _xlsx(client, auth, **params):
    r = client.get("/admin/knowledge-gaps/export.xlsx", params=params, headers=auth(HR))
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "attachment" in r.headers["content-disposition"]
    return r.content


def test_export_has_columns_filter_freeze_and_respects_filters(client, auth, db):
    ask(client, auth, OFF_TOPIC)
    client.post("/ai/feedback", json={"question": IN_KB, "answer_mode": "llm", "source_ids": ["KB-LEAVE-001"]},
                headers=auth(EMPLOYEE))
    ws = load_workbook(io.BytesIO(_xlsx(client, auth))).active
    header = [c.value for c in ws[1]]
    assert header == ["Date", "Question", "Language", "Feedback Type", "Matched Document",
                      "Response Status", "Review Status", "Category", "HR Comment"]
    assert ws.max_row == 3 and ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:I3"
    assert all(ws.column_dimensions[col].width >= 10 for col in "ABCDEFGHI")
    types = {ws.cell(row=r, column=4).value for r in (2, 3)}
    assert types == {"no_match", "negative_feedback"}

    only_feedback = load_workbook(io.BytesIO(_xlsx(client, auth, event_type="negative_feedback"))).active
    assert only_feedback.max_row == 2 and only_feedback.cell(row=2, column=5).value == "KB-LEAVE-001"

    # the chosen category is exported as "id: title"
    gap_id = events(db)[1].id
    client.patch(f"/admin/knowledge-gaps/{gap_id}", json={"category": "KB-LEAVE-002"}, headers=auth(HR))
    ws = load_workbook(io.BytesIO(_xlsx(client, auth, event_type="negative_feedback"))).active
    assert ws.cell(row=2, column=8).value == "KB-LEAVE-002: Порядок оформления и согласования отпуска"


def test_export_neutralises_formula_injection(client, auth, db):
    payloads = ['=HYPERLINK("http://evil.example","x") ДМС?', "+cmd|' /C calc'!A0 обед?", "@SUM(1+1) график?",
                "-2+3 столовая?"]
    for p in payloads:
        client.post("/ai/feedback", json={"question": p, "answer_mode": "llm"}, headers=auth(EMPLOYEE))
    gap_id = events(db)[0].id
    client.patch(f"/admin/knowledge-gaps/{gap_id}", json={"hr_comment": "=1+1"}, headers=auth(HR))

    content = _xlsx(client, auth)
    with zipfile.ZipFile(io.BytesIO(content)) as z:
        sheet = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert "<f>" not in sheet and "<f " not in sheet  # no formula is stored anywhere
    ws = load_workbook(io.BytesIO(content)).active
    questions = {ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)}
    assert questions == set(payloads)  # text preserved exactly
    for r in range(2, ws.max_row + 1):
        assert ws.cell(row=r, column=2).data_type == "s"


def test_questions_are_masked_and_export_has_no_personal_identifiers(client, auth, db):
    ask(client, auth, "Айдана тут. Мой email aidana.k@mail.example, телефон +7 701 555 01 01, ИИН 990101300123, есть ли ДМС?")
    [e] = events(db)
    assert "aidana.k@mail.example" not in e.question and "[EMAIL]" in e.question
    assert "555" not in e.question and "[ТЕЛЕФОН]" in e.question
    assert "990101300123" not in e.question and "[ИИН]" in e.question
    assert "Айдана" not in e.question and "[ИМЯ]" in e.question

    with zipfile.ZipFile(io.BytesIO(_xlsx(client, auth))) as z:
        text = " ".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.endswith(".xml"))
    for private in ("Касымова", "Айдана", "employee@janamen.example", "JP-0011", "aidana.k@mail.example", "990101300123"):
        assert private not in text
    assert not hasattr(KnowledgeGapEvent, "user_id")  # nothing to leak: who asked is never stored


# 10. the chat survives a broken journal ------------------------------------------

def test_chat_answers_even_if_gap_logging_fails(client, auth, db, monkeypatch):
    def broken(**_):
        raise RuntimeError("database is locked")
    monkeypatch.setattr(gaps, "KnowledgeGapEvent", broken)
    a = ask(client, auth, OFF_TOPIC)
    assert a["mode"] == "refusal" and a["escalation_suggested"] is True
    assert ask(client, auth, IN_KB)["grounded"] is True  # session still usable afterwards
    r = client.post("/ai/feedback", json={"question": IN_KB, "answer_mode": "llm"}, headers=auth(EMPLOYEE))
    assert r.status_code == 201 and r.json() == {"recorded": False}
