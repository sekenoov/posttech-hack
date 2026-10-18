"""Knowledge gap journal: employee feedback endpoint + HR review API and Excel export."""
import io
from datetime import date, datetime, time, timedelta
from typing import Annotated, Literal, Optional

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy import or_
from sqlalchemy.orm import Query as SAQuery, Session

from app.ai import gaps
from app.database import get_db
from app.models import KnowledgeDocument, KnowledgeGapEvent, Role, User
from app.permissions import get_current_user, not_found, require_roles
from app.schemas import UTCDateTime

router = APIRouter(tags=["knowledge-gaps"])
_hr = require_roles(Role.HR)

MAX_EXPORT_ROWS = 10000


class _Strict(BaseModel):
    model_config = ConfigDict(extra="forbid")


class FeedbackIn(_Strict):
    question: str = Field(min_length=2, max_length=1000)
    answer_mode: Literal["llm", "extractive", "personal", "refusal"]
    source_ids: list[Annotated[str, Field(pattern=r"^KB-[A-Z]+-\d{3}$")]] = Field(default_factory=list, max_length=5)


class GapOut(BaseModel):
    model_config = ConfigDict(from_attributes=True)

    id: int
    created_at: UTCDateTime
    event_type: str
    reason: str
    question: str
    language: str
    candidate_doc_id: Optional[str] = None
    answer_mode: str
    status: str
    category: Optional[str] = None
    hr_comment: Optional[str] = None


class GapPage(BaseModel):
    items: list[GapOut]
    total: int
    page: int
    page_size: int


class GapUpdate(_Strict):
    status: Optional[Literal["new", "reviewing", "resolved", "ignored"]] = None
    category: Optional[str] = Field(default=None, pattern=r"^KB-[A-Z]+-\d{3}$")  # knowledge document id; null clears
    hr_comment: Optional[str] = Field(default=None, max_length=2000)


# --- employee: «Не помогло» ---------------------------------------------------

@router.post("/ai/feedback", status_code=status.HTTP_201_CREATED)
def negative_feedback(body: FeedbackIn, _: User = Depends(get_current_user), db: Session = Depends(get_db)):
    """Anonymous: the event does not store who sent it."""
    event = gaps.log_feedback(db, body.question.strip(), body.answer_mode, body.source_ids)
    return {"recorded": event is not None}


# --- HR -----------------------------------------------------------------------

def _filtered(db: Session, q: Optional[str], status_: Optional[str], event_type: Optional[str],
              date_from: Optional[date], date_to: Optional[date]) -> SAQuery:
    query = db.query(KnowledgeGapEvent)
    if status_:
        query = query.filter(KnowledgeGapEvent.status == status_)
    if event_type:
        query = query.filter(KnowledgeGapEvent.event_type == event_type)
    if date_from:
        query = query.filter(KnowledgeGapEvent.created_at >= datetime.combine(date_from, time.min))
    if date_to:
        query = query.filter(KnowledgeGapEvent.created_at < datetime.combine(date_to + timedelta(days=1), time.min))
    if q and q.strip():
        like = f"%{q.strip()}%"
        query = query.filter(or_(KnowledgeGapEvent.question.ilike(like), KnowledgeGapEvent.hr_comment.ilike(like),
                                 KnowledgeGapEvent.category.ilike(like)))
    return query.order_by(KnowledgeGapEvent.created_at.desc(), KnowledgeGapEvent.id.desc())


class _Filters:
    def __init__(
        self,
        q: Optional[str] = Query(None, max_length=200),
        status_: Optional[Literal["new", "reviewing", "resolved", "ignored"]] = Query(None, alias="status"),
        event_type: Optional[Literal["no_match", "negative_feedback"]] = Query(None),
        date_from: Optional[date] = Query(None),
        date_to: Optional[date] = Query(None),
    ):
        self.args = (q, status_, event_type, date_from, date_to)


@router.get("/admin/knowledge-gaps", response_model=GapPage)
def list_gaps(
    filters: _Filters = Depends(),
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    _: User = Depends(_hr),
    db: Session = Depends(get_db),
):
    query = _filtered(db, *filters.args)
    total = query.count()
    items = query.offset((page - 1) * page_size).limit(page_size).all()
    return GapPage(items=[GapOut.model_validate(e) for e in items], total=total, page=page, page_size=page_size)


_COLUMNS = [
    ("Date", 20), ("Question", 60), ("Language", 10), ("Feedback Type", 18), ("Matched Document", 18),
    ("Response Status", 26), ("Review Status", 14), ("Category", 40), ("HR Comment", 40),
]
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "＝", "＋", "－", "＠")


def _put(ws, row: int, col: int, value):
    """Write a cell; user text that looks like a formula is stored as plain text."""
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        cell.data_type = "s"       # never written as <f>, so Excel will not evaluate it
        cell.quotePrefix = True    # and keeps it literal if the user edits the cell
    return cell


@router.get("/admin/knowledge-gaps/export.xlsx")
def export_gaps(filters: _Filters = Depends(), _: User = Depends(_hr), db: Session = Depends(get_db)):
    """Same filters as the list. No employee identifiers: none are stored."""
    rows = _filtered(db, *filters.args).limit(MAX_EXPORT_ROWS).all()
    titles = dict(db.query(KnowledgeDocument.id, KnowledgeDocument.title))
    wb = Workbook()
    ws = wb.active
    ws.title = "Knowledge gaps"
    header_fill = PatternFill("solid", fgColor="E3F6E3")
    for col, (title, width) in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        ws.column_dimensions[get_column_letter(col)].width = width
    for r, e in enumerate(rows, start=2):
        values = [
            e.created_at.strftime("%Y-%m-%d %H:%M UTC") if e.created_at else "",
            e.question, e.language, e.event_type, e.candidate_doc_id or "",
            f"{e.reason} ({e.answer_mode})", e.status,
            f"{e.category}: {titles.get(e.category, '')}".rstrip(": ") if e.category else "", e.hr_comment or "",
        ]
        for c, value in enumerate(values, start=1):
            _put(ws, r, c, value)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_COLUMNS))}{max(1, len(rows) + 1)}"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = f"janamen-knowledge-gaps-{date.today().isoformat()}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/admin/knowledge-gaps/{gap_id}", response_model=GapOut)
def get_gap(gap_id: int, _: User = Depends(_hr), db: Session = Depends(get_db)):
    event = db.get(KnowledgeGapEvent, gap_id)
    if event is None:
        raise not_found("Запись")
    return event


@router.patch("/admin/knowledge-gaps/{gap_id}", response_model=GapOut)
def update_gap(gap_id: int, body: GapUpdate, _: User = Depends(_hr), db: Session = Depends(get_db)):
    event = db.get(KnowledgeGapEvent, gap_id)
    if event is None:
        raise not_found("Запись")
    data = body.model_dump(exclude_unset=True)
    if "status" in data and data["status"]:
        event.status = data["status"]
    if "category" in data:
        if data["category"] and db.get(KnowledgeDocument, data["category"]) is None:
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Документ базы знаний не найден")
        event.category = data["category"] or None
    if "hr_comment" in data:
        event.hr_comment = (data["hr_comment"] or "").strip() or None
    db.commit()
    db.refresh(event)
    return event

